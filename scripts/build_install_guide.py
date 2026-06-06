from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "install-guide-assets"
VERSION = "0.1.3"
PDF_PATH = DOCS / "archive-timestamp-helper-install-guide-ja.pdf"
DOCX_PATH = DOCS / "archive-timestamp-helper-install-guide-ja.docx"
TXT_PATH = DOCS / "archive-timestamp-helper-install-guide-ja.txt"

FONT_REGULAR = Path("C:/Windows/Fonts/YuGothR.ttc")
FONT_MEDIUM = Path("C:/Windows/Fonts/YuGothM.ttc")
FONT_BOLD = Path("C:/Windows/Fonts/YuGothB.ttc")

PAGE_W = 1240
PAGE_H = 1754
MARGIN_X = 92
MARGIN_TOP = 76
INK = "#1f2937"
MUTED = "#64748b"
BLUE = "#2563eb"
BLUE_DARK = "#1d4ed8"
GREEN = "#16a34a"
RED = "#dc2626"
AMBER = "#d97706"
LINE = "#d8dee9"
FILL = "#f8fafc"


def font(size, weight="regular"):
    path = {
        "regular": FONT_REGULAR,
        "medium": FONT_MEDIUM,
        "bold": FONT_BOLD,
    }[weight]
    return ImageFont.truetype(str(path), size)


FONTS = {
    "title": font(42, "bold"),
    "subtitle": font(24, "medium"),
    "h1": font(30, "bold"),
    "h2": font(23, "bold"),
    "body": font(19, "regular"),
    "body_bold": font(19, "bold"),
    "small": font(16, "regular"),
    "small_bold": font(16, "bold"),
    "tiny": font(13, "regular"),
    "mono": font(17, "medium"),
    "label": font(18, "bold"),
    "badge": font(16, "bold"),
}


def text_size(draw, text, fnt):
    bbox = draw.textbbox((0, 0), text, font=fnt)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw, text, fnt, max_width):
    lines = []
    for para in text.splitlines():
        if not para:
            lines.append("")
            continue
        current = ""
        for char in para:
            trial = current + char
            if text_size(draw, trial, fnt)[0] <= max_width:
                current = trial
                continue
            if current:
                lines.append(current)
            current = char
        if current:
            lines.append(current)
    return lines


def draw_wrapped(draw, xy, text, fnt, fill=INK, max_width=900, line_gap=8):
    x, y = xy
    for line in wrap_text(draw, text, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += text_size(draw, line or " ", fnt)[1] + line_gap
    return y


def rounded(draw, box, fill, outline=None, width=1, radius=16):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def pill(draw, x, y, text, fill, color="#ffffff"):
    tw, th = text_size(draw, text, FONTS["badge"])
    box = (x, y, x + tw + 28, y + th + 18)
    rounded(draw, box, fill=fill, radius=20)
    draw.text((x + 14, y + 8), text, font=FONTS["badge"], fill=color)
    return box[2]


def marker(draw, x, y, number, color=RED):
    draw.ellipse((x, y, x + 44, y + 44), fill=color)
    tw, th = text_size(draw, str(number), FONTS["label"])
    draw.text((x + 22 - tw / 2, y + 20 - th / 2), str(number), font=FONTS["label"], fill="#ffffff")


def page_base(page_no, title):
    img = Image.new("RGB", (PAGE_W, PAGE_H), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, PAGE_W, 18), fill=BLUE)
    draw.text((MARGIN_X, MARGIN_TOP), title, font=FONTS["title"], fill="#0f172a")
    draw.line((MARGIN_X, PAGE_H - 72, PAGE_W - MARGIN_X, PAGE_H - 72), fill="#e2e8f0", width=2)
    footer = f"Archive Timestamp Helper Beta v{VERSION}  |  {page_no}"
    draw.text((MARGIN_X, PAGE_H - 50), footer, font=FONTS["tiny"], fill=MUTED)
    return img, draw


def draw_callout(draw, x, y, w, title, body, color=BLUE):
    rounded(draw, (x, y, x + w, y + 132), fill="#eff6ff", outline="#bfdbfe", radius=18, width=2)
    draw.text((x + 24, y + 18), title, font=FONTS["h2"], fill=color)
    draw_wrapped(draw, (x + 24, y + 58), body, FONTS["body"], fill=INK, max_width=w - 48, line_gap=7)


def make_extract_figure(path):
    img = Image.new("RGB", (980, 460), "#f8fafc")
    d = ImageDraw.Draw(img)
    rounded(d, (30, 30, 950, 430), fill="#ffffff", outline="#cbd5e1", radius=20, width=3)
    d.text((62, 62), "ZIPを展開してから読み込みます", font=FONTS["h2"], fill="#0f172a")
    items = [
        ("1", "ZIPをダウンロード", f"archive-timestamp-helper-beta-{VERSION}.zip"),
        ("2", "右クリック → すべて展開", "ZIPのままでは読み込めません"),
        ("3", "展開したフォルダを選ぶ", "中に manifest.json があるフォルダ"),
    ]
    x = 72
    for no, title, note in items:
        marker(d, x, 150, no, BLUE)
        rounded(d, (x + 60, 138, x + 280, 310), fill="#f1f5f9", outline="#dbe3ed", radius=16, width=2)
        d.text((x + 82, 166), title, font=FONTS["small_bold"], fill=INK)
        draw_wrapped(d, (x + 82, 208), note, FONTS["small"], fill=MUTED, max_width=176, line_gap=5)
        if no != "3":
            d.line((x + 295, 225, x + 340, 225), fill=BLUE, width=5)
            d.polygon([(x + 340, 225), (x + 323, 212), (x + 323, 238)], fill=BLUE)
        x += 300
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def draw_browser_frame(d, x, y, w, h, browser, url, accent):
    rounded(d, (x, y, x + w, y + h), fill="#ffffff", outline="#cbd5e1", radius=18, width=3)
    d.rounded_rectangle((x, y, x + w, y + 62), radius=18, fill="#eef2f7")
    d.rectangle((x, y + 36, x + w, y + 62), fill="#eef2f7")
    for i, color in enumerate(["#ef4444", "#f59e0b", "#22c55e"]):
        d.ellipse((x + 22 + i * 24, y + 22, x + 36 + i * 24, y + 36), fill=color)
    rounded(d, (x + 120, y + 16, x + w - 30, y + 46), fill="#ffffff", outline="#d1d5db", radius=14)
    d.text((x + 138, y + 20), url, font=FONTS["tiny"], fill="#334155")
    d.text((x + 28, y + 86), browser, font=FONTS["h2"], fill="#0f172a")
    d.text((x + 28, y + 130), "拡張機能", font=FONTS["h1"], fill="#0f172a")
    return x + 28, y + 172, w - 56


def make_chrome_figure(path):
    img = Image.new("RGB", (980, 620), "#f8fafc")
    d = ImageDraw.Draw(img)
    cx, cy, cw = draw_browser_frame(d, 22, 22, 936, 576, "Chrome", "chrome://extensions", BLUE)
    d.text((cx, cy), "デベロッパー モード", font=FONTS["small_bold"], fill=INK)
    rounded(d, (cx + 800, cy - 4, cx + 882, cy + 34), fill=BLUE, radius=20)
    d.ellipse((cx + 842, cy, cx + 878, cy + 34), fill="#ffffff")
    d.rectangle((cx, cy + 70, cx + 278, cy + 124), fill="#eef2ff", outline=BLUE, width=4)
    d.text((cx + 20, cy + 88), "パッケージ化されていない拡張機能を読み込む", font=FONTS["tiny"], fill=BLUE_DARK)
    rounded(d, (cx, cy + 176, cx + 360, cy + 360), fill="#ffffff", outline="#dbe3ed", radius=16, width=2)
    d.text((cx + 24, cy + 202), "Archive Timestamp Helper Beta", font=FONTS["small_bold"], fill=INK)
    d.text((cx + 24, cy + 240), f"バージョン {VERSION}", font=FONTS["small"], fill=MUTED)
    d.text((cx + 24, cy + 284), "有効", font=FONTS["small_bold"], fill=GREEN)
    marker(d, cx + 790, cy - 60, 1, RED)
    d.rectangle((cx + 790, cy - 4, cx + 890, cy + 40), outline=RED, width=5)
    marker(d, cx + 300, cy + 70, 2, RED)
    d.rectangle((cx - 4, cy + 66, cx + 282, cy + 128), outline=RED, width=5)
    marker(d, cx + 382, cy + 190, 3, RED)
    d.rectangle((cx - 4, cy + 172, cx + 364, cy + 364), outline=RED, width=5)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def make_edge_figure(path):
    img = Image.new("RGB", (980, 620), "#f8fafc")
    d = ImageDraw.Draw(img)
    cx, cy, cw = draw_browser_frame(d, 22, 22, 936, 576, "Microsoft Edge", "edge://extensions", "#0f9f8f")
    d.rectangle((cx, cy - 10, cx + 220, cy + 386), fill="#f1f5f9", outline="#dbe3ed")
    d.text((cx + 20, cy + 18), "インストール済みの拡張機能", font=FONTS["tiny"], fill=INK)
    d.text((cx + 20, cy + 322), "開発者モード", font=FONTS["small_bold"], fill=INK)
    rounded(d, (cx + 144, cy + 316, cx + 204, cy + 348), fill="#0f9f8f", radius=18)
    d.ellipse((cx + 176, cy + 319, cx + 201, cy + 345), fill="#ffffff")
    d.rectangle((cx + 260, cy + 28, cx + 458, cy + 82), fill="#ecfeff", outline="#0f9f8f", width=4)
    d.text((cx + 280, cy + 46), "展開して読み込み", font=FONTS["small_bold"], fill="#0f766e")
    rounded(d, (cx + 260, cy + 136, cx + 760, cy + 318), fill="#ffffff", outline="#dbe3ed", radius=16, width=2)
    d.text((cx + 284, cy + 162), "Archive Timestamp Helper Beta", font=FONTS["small_bold"], fill=INK)
    d.text((cx + 284, cy + 200), f"バージョン {VERSION}", font=FONTS["small"], fill=MUTED)
    d.text((cx + 284, cy + 244), "有効", font=FONTS["small_bold"], fill=GREEN)
    marker(d, cx + 224, cy + 314, 1, RED)
    d.rectangle((cx + 12, cy + 306, cx + 210, cy + 356), outline=RED, width=5)
    marker(d, cx + 470, cy + 28, 2, RED)
    d.rectangle((cx + 256, cy + 24, cx + 462, cy + 86), outline=RED, width=5)
    marker(d, cx + 782, cy + 150, 3, RED)
    d.rectangle((cx + 256, cy + 132, cx + 764, cy + 322), outline=RED, width=5)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def make_panel_figure(path):
    img = Image.new("RGB", (980, 520), "#eef2f7")
    d = ImageDraw.Draw(img)
    rounded(d, (40, 36, 940, 484), fill="#111827", outline="#0f172a", radius=24, width=2)
    d.rectangle((64, 66, 690, 434), fill="#1f2937", outline="#374151")
    d.text((92, 96), "まさおの家 ライブアーカイブ", font=FONTS["h2"], fill="#ffffff")
    d.text((92, 140), "動画ページ", font=FONTS["body"], fill="#cbd5e1")
    rounded(d, (620, 92, 906, 410), fill="#ffffff", outline="#d1d5db", radius=18, width=3)
    d.text((646, 116), "Timestamp Helper", font=FONTS["small_bold"], fill=INK)
    d.text((646, 146), "-3秒 / Alt+キー", font=FONTS["tiny"], fill=MUTED)
    labels = [("+ 時刻", "#f1f5f9"), ("ゴロン", "#eff6ff"), ("あくび", "#eff6ff"), ("ちもしー", "#eff6ff"), ("コピー", "#dcfce7")]
    yy = 184
    for label, fill in labels:
        rounded(d, (646, yy, 878, yy + 42), fill=fill, outline="#cbd5e1", radius=10)
        d.text((664, yy + 10), label, font=FONTS["small_bold"], fill=INK)
        yy += 52
    marker(d, 886, 104, 1, RED)
    d.text((712, 432), "右下に出たらOK", font=FONTS["small_bold"], fill="#ffffff")
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def image_block(page, draw, img_path, x, y, w):
    src = Image.open(img_path).convert("RGB")
    h = int(src.height * (w / src.width))
    src = src.resize((w, h), Image.Resampling.LANCZOS)
    draw.rounded_rectangle((x - 4, y - 4, x + w + 4, y + h + 4), radius=18, fill="#e2e8f0")
    page.paste(src, (x, y))
    return y + h


def draw_steps(draw, x, y, steps, max_width=920):
    for idx, (title, body) in enumerate(steps, 1):
        marker(draw, x, y, idx, BLUE)
        draw.text((x + 62, y + 2), title, font=FONTS["body_bold"], fill=INK)
        y = draw_wrapped(draw, (x + 62, y + 32), body, FONTS["body"], fill=INK, max_width=max_width - 62, line_gap=7)
        y += 18
    return y


def make_pages():
    extract = ASSETS / "zip-extract-flow.png"
    chrome = ASSETS / "chrome-extension-page.png"
    edge = ASSETS / "edge-extension-page.png"
    panel = ASSETS / "helper-panel.png"
    make_extract_figure(extract)
    make_chrome_figure(chrome)
    make_edge_figure(edge)
    make_panel_figure(panel)

    pages = []

    img, d = page_base(1, "インストールガイド")
    d.text((MARGIN_X, 142), "Archive Timestamp Helper Beta", font=FONTS["subtitle"], fill=BLUE_DARK)
    pill(d, MARGIN_X, 188, f"手動インストール版 v{VERSION}", BLUE)
    d.text((MARGIN_X, 250), "まさおの家のYouTubeライブアーカイブで、タイムスタンプ作成を手伝う試作品です。", font=FONTS["body"], fill=INK)
    draw_callout(
        d,
        MARGIN_X,
        302,
        PAGE_W - MARGIN_X * 2,
        "先に大事なこと",
        "ZIPファイルは必ず「すべて展開」してから読み込みます。ZIPのまま選んでも、Chrome/Edgeは拡張機能として読み込めません。",
        BLUE,
    )
    image_block(img, d, extract, MARGIN_X, 462, PAGE_W - MARGIN_X * 2)
    d.text((MARGIN_X, 965), "準備するもの", font=FONTS["h1"], fill="#0f172a")
    y = draw_steps(d, MARGIN_X, 1030, [
        ("ZIPファイルをダウンロード", f"配布ページから archive-timestamp-helper-beta-{VERSION}.zip を保存します。"),
        ("ZIPを右クリックして展開", "Windowsなら「すべて展開」を選びます。展開先はデスクトップやダウンロードフォルダで大丈夫です。"),
        ("manifest.json があるフォルダを確認", "読み込み時に選ぶのは、manifest.json が入っているフォルダです。"),
    ])
    draw_callout(d, MARGIN_X, y + 12, PAGE_W - MARGIN_X * 2, "この拡張機能がすること", "コメントの自動投稿はしません。記録したタイムスタンプをコピーして、人が確認してからコメント欄に貼るための道具です。", GREEN)
    pages.append(img)

    img, d = page_base(2, "Chromeで入れる")
    image_block(img, d, chrome, MARGIN_X, 170, PAGE_W - MARGIN_X * 2)
    d.text((MARGIN_X, 875), "手順", font=FONTS["h1"], fill="#0f172a")
    draw_steps(d, MARGIN_X, 940, [
        ("chrome://extensions を開く", "Chromeのアドレスバーに chrome://extensions と入力してEnterを押します。"),
        ("デベロッパーモードをオン", "右上の「デベロッパーモード」をオンにします。"),
        ("読み込みボタンを押す", "「パッケージ化されていない拡張機能を読み込む」を押します。"),
        ("展開したフォルダを選ぶ", "ZIPを展開したフォルダを選択します。Archive Timestamp Helper Beta が表示されたら完了です。"),
    ])
    draw_callout(d, MARGIN_X, 1450, PAGE_W - MARGIN_X * 2, "Chromeで更新するとき", "同じフォルダの中身を新しい版に置き換えた場合は、拡張機能ページで「再読み込み」を押すだけで更新できます。", AMBER)
    pages.append(img)

    img, d = page_base(3, "Edgeで入れる")
    image_block(img, d, edge, MARGIN_X, 170, PAGE_W - MARGIN_X * 2)
    d.text((MARGIN_X, 875), "手順", font=FONTS["h1"], fill="#0f172a")
    draw_steps(d, MARGIN_X, 940, [
        ("edge://extensions を開く", "Edgeのアドレスバーに edge://extensions と入力してEnterを押します。"),
        ("開発者モードをオン", "左側メニュー、または画面内の「開発者モード」をオンにします。"),
        ("読み込みボタンを押す", "「展開して読み込み」または「パッケージ化されていない拡張機能を読み込む」を押します。"),
        ("展開したフォルダを選ぶ", "ZIPを展開したフォルダを選択します。Archive Timestamp Helper Beta が表示されたら完了です。"),
    ])
    draw_callout(d, MARGIN_X, 1450, PAGE_W - MARGIN_X * 2, "Edgeで警告が出たら", "手動で入れた拡張機能なので、Edgeが確認表示を出すことがあります。自分で入れたものだと分かる場合だけ有効にしてください。", AMBER)
    pages.append(img)

    img, d = page_base(4, "使い方と困った時")
    image_block(img, d, panel, MARGIN_X, 170, PAGE_W - MARGIN_X * 2)
    d.text((MARGIN_X, 760), "使い方", font=FONTS["h1"], fill="#0f172a")
    draw_steps(d, MARGIN_X, 825, [
        ("まさおの家のライブアーカイブを開く", "対象のアーカイブだけで、右下に Timestamp Helper のパネルが出ます。"),
        ("ボタンかショートカットで記録", "Alt+G はゴロン、Alt+A はあくび、Alt+C はちもしー、Alt+T は時刻だけ、Alt+Z は直前取り消しです。"),
        ("コピーしてコメント欄へ", "最後に「コピー」を押すと、コメント欄に貼り付けやすい形でコピーされます。"),
    ])
    d.text((MARGIN_X, 1220), "パネルが出ない時", font=FONTS["h1"], fill="#0f172a")
    y = draw_steps(d, MARGIN_X, 1285, [
        ("対象を確認", "他チャンネル、通常動画、配信中のライブでは表示されません。まさおの家のライブアーカイブで確認してください。"),
        ("ページと拡張機能を再読み込み", "YouTubeページを再読み込みし、拡張機能ページでも「再読み込み」を押します。"),
        ("ZIPではなく展開フォルダ", "読み込み先がZIPファイルではなく、manifest.json があるフォルダになっているか確認します。"),
    ], max_width=980)
    pages.append(img)

    return pages


def save_pdf(pages):
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4)
    page_w, page_h = A4
    for page in pages:
        temp = ASSETS / "_pdf_page.png"
        page.save(temp)
        c.drawImage(ImageReader(str(temp)), 0, 0, width=page_w, height=page_h)
        c.showPage()
    c.save()
    temp.unlink(missing_ok=True)


def save_docx():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Yu Gothic"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Archive Timestamp Helper Beta インストールガイド")
    run.bold = True
    run.font.size = Pt(20)

    sub = document.add_paragraph(f"Chrome / Edge 手動インストール版 v{VERSION}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sections = [
        ("準備", [
            "ZIPファイルは必ず「すべて展開」してから読み込みます。",
            "読み込み時に選ぶのは、manifest.json が入っているフォルダです。",
            "コメントの自動投稿はしません。コピーして人が確認してから貼り付ける道具です。",
        ], ASSETS / "zip-extract-flow.png"),
        ("Chromeで入れる", [
            "Chromeのアドレスバーに chrome://extensions と入力します。",
            "右上の「デベロッパーモード」をオンにします。",
            "「パッケージ化されていない拡張機能を読み込む」を押します。",
            "ZIPを展開したフォルダを選択します。",
        ], ASSETS / "chrome-extension-page.png"),
        ("Edgeで入れる", [
            "Edgeのアドレスバーに edge://extensions と入力します。",
            "「開発者モード」をオンにします。",
            "「展開して読み込み」または「パッケージ化されていない拡張機能を読み込む」を押します。",
            "ZIPを展開したフォルダを選択します。",
        ], ASSETS / "edge-extension-page.png"),
        ("使い方と困った時", [
            "まさおの家のライブアーカイブだけで、右下にパネルが出ます。",
            "Alt+G: ゴロン / Alt+A: あくび / Alt+C: ちもしー / Alt+T: 時刻だけ / Alt+Z: 直前取り消し",
            "パネルが出ない時は、対象アーカイブか、ページ再読み込み、拡張機能の再読み込み、展開フォルダを確認します。",
        ], ASSETS / "helper-panel.png"),
    ]

    for heading, bullets, image in sections:
        document.add_heading(heading, level=1)
        for item in bullets:
            document.add_paragraph(item, style="List Bullet")
        document.add_picture(str(image), width=Inches(6.8))

    document.save(DOCX_PATH)


def save_txt():
    text = f"""Archive Timestamp Helper Beta インストールガイド
Chrome / Edge 手動インストール版 v{VERSION}

1. まずZIPを展開
- archive-timestamp-helper-beta-{VERSION}.zip をダウンロードします。
- ZIPを右クリックして「すべて展開」を選びます。
- 読み込み時に選ぶのは、manifest.json が入っている展開済みフォルダです。

2. Chromeで入れる
- Chromeで chrome://extensions を開きます。
- 右上の「デベロッパーモード」をオンにします。
- 「パッケージ化されていない拡張機能を読み込む」を押します。
- ZIPを展開したフォルダを選びます。

3. Edgeで入れる
- Edgeで edge://extensions を開きます。
- 「開発者モード」をオンにします。
- 「展開して読み込み」または「パッケージ化されていない拡張機能を読み込む」を押します。
- ZIPを展開したフォルダを選びます。

4. 使い方
- まさおの家のライブアーカイブを開くと、右下にパネルが出ます。
- Alt+G: ゴロン
- Alt+A: あくび
- Alt+C: ちもしー
- Alt+T: 時刻だけ追加
- Alt+Z: 直前の追加を取り消し
- 「コピー」を押すと、コメント欄に貼り付けやすい形でコピーされます。

5. 注意
- コメントの自動投稿はしません。
- 他チャンネル、通常動画、配信中のライブでは表示されません。
- 記録した内容は自分のブラウザ内に保存されます。
"""
    TXT_PATH.write_text(text, encoding="utf-8")


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    pages = make_pages()
    for index, page in enumerate(pages, 1):
        page.save(ASSETS / f"page-{index}.png")
    save_pdf(pages)
    save_docx()
    save_txt()
    print(PDF_PATH)
    print(DOCX_PATH)
    print(TXT_PATH)


if __name__ == "__main__":
    main()
